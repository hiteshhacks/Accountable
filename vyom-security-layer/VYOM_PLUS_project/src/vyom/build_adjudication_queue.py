"""Build the human adjudication queue for the disputed reviewer records.

Source of truth: data/synthetic/VYOM_plus_reviewer_comparison_adjudication.xlsx,
sheet Adjudication_Queue (every record with needs_adjudication = True). Record IDs
and transaction fields come from the human_review_queue sheet of the synthetic
dataset, matched by review case number (HR-0001 = queue row 1).

Nothing is decided here. Reviewer labels, notes and fields are copied verbatim;
the added columns only describe the disagreement, quote the evidence each
reviewer cited, list factual field conflicts or gaps, and set a review priority.
Final labels and rationales are left blank with status PENDING.

Source workbooks are read only. The output is refused if it already exists.
"""

from pathlib import Path
import argparse
import hashlib
import re

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation


ROOT = Path(__file__).resolve().parents[2]
COMPARISON = ROOT / "data/synthetic/VYOM_plus_reviewer_comparison_adjudication.xlsx"
DATASET = ROOT / "data/synthetic/synthetic_voucher_dataset_v5_review_ready.xlsx"
OUTPUT = ROOT / "reports/vyom_adjudication_queue.xlsx"

NON_FINAL = "NEEDS_ADJUDICATION"
STATUSES = ["PENDING", "ADJUDICATED", "NEEDS_POLICY_DECISION", "INSUFFICIENT_EVIDENCE"]

TRANSACTION_FIELDS = [
    "seller_supplier", "buyer_customer", "invoice_number", "transaction_date", "item_description",
    "transaction_narration", "quantity", "unit_price", "taxable_value", "gst_rate_percent", "gst_amount",
    "discount_amount", "freight_amount", "currency", "payment_method", "payment_status", "order_reference",
    "return_reason", "debit_credit_info", "payroll_period", "employee_count", "attendance_days",
    "import_export_reference", "customs_duty_amount", "warehouse_from", "warehouse_to", "movement_reason",
    "document_type", "gross_pay_amount", "deductions_amount", "net_pay_amount",
]
RELEVANT_FIELDS = [
    "transaction_narration", "item_description", "payment_status", "return_reason", "debit_credit_info",
    "order_reference", "movement_reason", "warehouse_from", "warehouse_to", "document_type",
    "seller_supplier", "buyer_customer", "payment_method", "taxable_value",
]

# Definition issues (detailed in reports/vyom_label_definition_gaps.md).
DEFINITION_ISSUES = {
    "DEF-01": "Rejection In vs Rejection Out: which party rejects, and the direction naming convention",
    "DEF-02": "Physical rejection vs financial return (Rejection In/Out vs Debit/Credit Note)",
    "DEF-03": "Evidence required for a return / credit note (link to the original sale)",
    "DEF-04": "Payment vs Purchase Return / Debit Note (settlement vs adjustment of a payable)",
    "DEF-05": "Payment vs Receipt vs Contra (same-entity accounts vs external parties)",
    "DEF-06": "Payment vs Expense vs Purchase (cash settlement vs cost recognition; consumables policy)",
    "DEF-07": "Advance / Prepayment direction (paid vs received) and relation to Payment/Receipt",
    "DEF-08": "Material In / Material Out scope (internal movement vs subcontracting) and rejection fields",
    "DEF-09": "Job Work In Order vs Job Work Out Order (provisional definitions)",
    "DEF-10": "Other / Miscellaneous: voucher class or review status",
    "DEF-11": "Purchase Return vs Sales Return direction",
}

# Disagreement pattern -> (definition issues, candidate categories to consider, decision question).
PATTERNS = {
    "Rejection In vs Rejection Out": (
        ["DEF-01"], ["Rejection In", "Rejection Out"],
        "Who rejected the goods: the entity on goods received from a supplier, or a customer on goods the entity "
        "delivered? Which naming convention does the target accounting software use (DEF-01)?"),
    "Purchase Return / Debit Note vs Rejection In": (
        ["DEF-02", "DEF-01"], ["Rejection In", "Purchase Return / Debit Note"],
        "Had the supplier purchase already been booked, and is this record the physical rejection or the "
        "financial adjustment (debit note) of that purchase?"),
    "Payment vs Purchase Return / Debit Note": (
        ["DEF-04"], ["Purchase Return / Debit Note", "Payment"],
        "Does the entry move cash or bank funds, or does it only reduce a supplier payable (debit note)?"),
    "Contra vs NEEDS_ADJUDICATION": (
        ["DEF-05"], ["Contra", "Payment", "Receipt"],
        "Are both accounts named and owned by the same entity? Which accounts are debited and credited?"),
    "Advance / Prepayment vs NEEDS_ADJUDICATION": (
        ["DEF-07"], ["Advance / Prepayment", "Payment", "Receipt"],
        "Was the advance paid by the entity or received from a customer, and is direction part of the category?"),
    "NEEDS_ADJUDICATION vs Sales Return / Credit Note": (
        ["DEF-03"], ["Sales Return / Credit Note", "Rejection Out"],
        "Is a link to the original sale invoice (or a credit-note number) required before labelling a return?"),
    "Journal vs NEEDS_ADJUDICATION": (
        [], ["Journal"],
        "Does a non-cash adjustment with status 'Pending approval' qualify as a Journal voucher, or must it be posted?"),
    "NEEDS_ADJUDICATION vs Rejection Out": (
        ["DEF-01", "DEF-02"], ["Rejection Out", "Rejection In", "Sales Return / Credit Note"],
        "Who rejected the goods, and has a credit note been booked? Do conflicting status fields change the answer?"),
    "Material Out vs Rejection In": (
        ["DEF-08", "DEF-01"], ["Material Out", "Rejection In"],
        "Is the primary event an internal material issue or a quality rejection? Which field is authoritative when "
        "document_type and movement_reason disagree?"),
    "Material In vs Rejection In": (
        ["DEF-08", "DEF-01"], ["Material In", "Rejection In"],
        "Is the primary event an internal material receipt or a quality rejection? Which field is authoritative when "
        "document_type and movement_reason disagree?"),
    "Sales vs Sales Return / Credit Note": (
        [], ["Sales Return / Credit Note", "Sales"],
        "The guideline excludes returns/credit notes from Sales; confirm whether this revenue reversal is a return."),
    "Rejection Out vs Sales Return / Credit Note": (
        ["DEF-02"], ["Sales Return / Credit Note", "Rejection Out"],
        "Is this the physical customer rejection or the financial credit to the customer balance?"),
    "NEEDS_ADJUDICATION vs Payment": (
        ["DEF-06"], ["Payment", "Expense", "Purchase"],
        "Is the primary event settlement of an existing payable (Payment) or recognition of a cost (Expense/Purchase)?"),
    "Expense vs NEEDS_ADJUDICATION": (
        ["DEF-06"], ["Expense", "Payment"],
        "Is the primary event settlement of an existing payable (Payment) or recognition of a cost (Expense)?"),
    "NEEDS_ADJUDICATION vs Purchase": (
        ["DEF-06"], ["Purchase", "Expense"],
        "Are office supplies goods acquisition (Purchase) or an operating cost (Expense) under the chart-of-accounts policy?"),
    "NEEDS_ADJUDICATION vs Rejection In": (
        ["DEF-01", "DEF-02"], ["Rejection In", "Purchase Return / Debit Note"],
        "Who rejected the goods, and has a debit note been booked?"),
    "both: Rejection In": (
        ["DEF-01"], ["Rejection In", "Rejection Out"],
        "Both reviewers chose Rejection In but flagged the record; confirm the label once DEF-01 is decided and the "
        "status field ('Accepted for review') is explained."),
}
BOTH_ABSTAINED_OTHER = (
    ["DEF-10"], ["Other / Miscellaneous (if kept as a class)", "Review status (if Other means 'send to review')"],
    "Should a record whose own narration says the event is unclear be labelled Other / Miscellaneous, or carry a "
    "review status with no voucher class (DEF-10)?")
BOTH_ABSTAINED = (
    [], [],
    "Both reviewers found no supported category; decide whether the record has enough evidence for any label.")


def disagreement_pair(r1: str, r2: str) -> str:
    if r1 == r2:
        return f"both: {r1}"
    return " vs ".join(sorted([r1, r2]))


def disagreement_kind(r1: str, r2: str, flagged: bool) -> str:
    if r1 == r2 == NON_FINAL:
        return "BOTH_ABSTAINED"
    if NON_FINAL in (r1, r2):
        return "ONE_REVIEWER_ABSTAINED"
    if r1 != r2:
        return "CONCRETE_LABEL_CONFLICT"
    return "SAME_LABEL_WITH_QUALITY_FLAG" if flagged else "SAME_LABEL"


def _text(value) -> str:
    return "" if pd.isna(value) else str(value).strip()


def narration_core(text: str) -> str:
    """Narration without the generated reference/date sentence."""
    m = re.search(r"TXN-\d+", text)
    return text[:text.rfind(".", 0, m.start()) + 1].strip() if m else text


ACCOUNT_DEPENDENT = {"Payment", "Receipt", "Contra", "Journal", "Purchase Return / Debit Note",
                     "Advance / Prepayment", "Expense", "Purchase"}


def field_findings(row: pd.Series, candidates=()) -> list:
    """Factual conflicts or gaps between fields. Never proposes a label."""
    found = []
    narration = _text(row.get("transaction_narration")).lower()
    status = _text(row.get("payment_status"))
    movement = _text(row.get("movement_reason"))
    doc_type = _text(row.get("document_type"))
    dci = _text(row.get("debit_credit_info"))
    item = _text(row.get("item_description")).lower()
    parties = [p for p in (_text(row.get("seller_supplier")), _text(row.get("buyer_customer"))) if p]

    if re.search(r"unclear|unresolved|does not support|manual review|held for review|should not be forced", narration):
        found.append("Narration itself states the event is unclear / held for review")
    if "customer" in narration and status == "Returned to supplier":
        found.append("Narration attributes the rejection to a customer, but payment_status = 'Returned to supplier'")
    if doc_type.startswith("Material") and movement == "Quality rejection":
        found.append(f"document_type = '{doc_type}' but movement_reason = 'Quality rejection'")
    if "reject" in narration and movement and movement != "Quality rejection":
        found.append(f"Narration describes a rejection, but movement_reason = '{movement}'")
    if re.search(r"own cash or bank|same entity|own-account", narration) and parties:
        if len(set(parties)) > 1:
            found.append(f"Narration describes an own-account transfer, but party fields name two different parties: "
                         f"{', '.join(parties)}")
        else:
            found.append(f"Narration describes an own-account transfer; party fields name only '{parties[0]}' "
                         f"(confirm this is the entity itself)")
    if "paid or received" in narration:
        found.append("Narration does not say whether the advance was paid or received")
    if "advance received" in item and status in ("Prepaid", "Paid in advance"):
        found.append(f"item_description = '{_text(row.get('item_description'))}' but payment_status = '{status}'")
    if dci and not re.search(r"cash|bank", dci, re.I) and re.search(r"payable|receivable", dci, re.I):
        found.append(f"debit_credit_info names no cash/bank account: '{dci}'")
    if not dci and ACCOUNT_DEPENDENT & set(candidates):
        found.append("debit_credit_info is empty (accounts debited/credited not recorded)")
    return found


def cited_evidence(row: pd.Series, reviewer: int) -> str:
    """Quote the evidence a reviewer cited, with values taken from the record itself."""
    label = _text(row[f"reviewer_{reviewer}_label"])
    conf = row[f"reviewer_{reviewer}_confidence_0_to_3"]
    flag = _text(row[f"reviewer_{reviewer}_data_quality_flag"]) or "none recorded"
    note = _text(row[f"reviewer_{reviewer}_notes"])
    raw = _text(row[f"reviewer_{reviewer}_evidence_fields"])
    if reviewer == 1:
        names = [n.strip() for n in raw.split(",") if n.strip()]
        quoted = "; ".join(
            f"{n}={narration_core(_text(row.get(n))) if n == 'transaction_narration' else (_text(row.get(n)) or '(empty)')}"
            for n in names)
    else:
        quoted = raw
    stance = "abstained" if label == NON_FINAL else f"chose {label}"
    return f"Reviewer {reviewer} {stance} (confidence {conf}/3, quality flag: {flag}). Note: {note} Cited: {quoted}"


def assign_priority(kind: str, defs: list, findings: list, is_other: bool):
    """Return (priority, reason). Criteria are listed in the README sheet and the summary."""
    if is_other:
        return "HIGH", "Both reviewers abstained on a record that decides the Other / Miscellaneous policy (DEF-10)"
    if kind == "CONCRETE_LABEL_CONFLICT" and defs:
        return "HIGH", f"Two concrete labels conflict on an open definition ({', '.join(defs)})"
    if kind == "CONCRETE_LABEL_CONFLICT":
        return "MEDIUM", "Two concrete labels conflict; the definition itself is not disputed"
    if kind == "SAME_LABEL_WITH_QUALITY_FLAG":
        return "MEDIUM", f"Reviewers agree but a quality flag was raised; label depends on {', '.join(defs) or 'record check'}"
    if defs:
        return "MEDIUM", f"One reviewer abstained on an open definition ({', '.join(defs)})"
    if kind == "BOTH_ABSTAINED":
        return "MEDIUM", "Both reviewers abstained; decide whether any label is supported"
    return "LOW", "One reviewer abstained; no open definition issue applies (see Field Conflicts or Gaps)"


def load_sources():
    comparison_sheets = pd.read_excel(COMPARISON, sheet_name=None)
    queue_src = comparison_sheets["Adjudication_Queue"]
    consensus = comparison_sheets["Provisional_Consensus"]
    review_queue = pd.read_excel(DATASET, sheet_name="human_review_queue")
    return comparison_sheets, queue_src, consensus, review_queue


def build_queue():
    """Return the queue and supporting tables as DataFrames (nothing written)."""
    sheets, src, consensus, review_queue = load_sources()
    disputed = src[src["needs_adjudication"] == True].copy()  # noqa: E712 - explicit boolean column
    if len(disputed) != len(src):
        raise ValueError("Adjudication_Queue contains rows not flagged needs_adjudication")
    disputed["queue_row"] = disputed["review_case_id"].str.extract(r"HR-(\d+)")[0].astype(int) - 1
    linked = review_queue.iloc[disputed["queue_row"]].reset_index(drop=True)
    disputed = disputed.reset_index().rename(columns={"index": "source_row"})
    shared = [c for c in TRANSACTION_FIELDS if c in disputed.columns]
    if not (disputed[shared].astype(str).values == linked[shared].astype(str).values).all():
        raise ValueError("Adjudication_Queue fields do not match human_review_queue rows")

    rows = []
    for i, r in disputed.iterrows():
        r1, r2 = _text(r["reviewer_1_label"]), _text(r["reviewer_2_label"])
        flagged = bool(_text(r["reviewer_1_data_quality_flag"])) or _text(r["reviewer_2_data_quality_flag"]) not in ("", "NONE")
        kind = disagreement_kind(r1, r2, flagged)
        pair = disagreement_pair(r1, r2)
        unclear = bool(re.search(r"unclear|unresolved|does not support|manual review|held for review|should not be forced",
                                 _text(r.get("transaction_narration")).lower()))
        is_other = kind == "BOTH_ABSTAINED" and unclear
        defs, candidates, question = PATTERNS.get(pair, BOTH_ABSTAINED_OTHER if is_other else BOTH_ABSTAINED)
        findings = field_findings(r, candidates + [r1, r2])
        if kind == "BOTH_ABSTAINED" and not is_other:
            defs = []
        priority, reason = assign_priority(kind, defs, findings, is_other)
        relevant = {f: _text(r.get(f)) for f in RELEVANT_FIELDS}
        relevant["transaction_narration"] = narration_core(relevant["transaction_narration"])
        gaps = [f for f in findings]
        rows.append({
            "Review Case ID": r["review_case_id"],
            "Record ID": linked.loc[i, "record_id"],
            "Source Record ID": linked.loc[i, "source_record_id"],
            "Source Reference": (f"{COMPARISON.name} / Adjudication_Queue Excel row {int(r['source_row']) + 2}; "
                                 f"{DATASET.name} / human_review_queue Excel row {int(r['queue_row']) + 2}"),
            "Reviewer 1 Label": r1,
            "Reviewer 2 Label": r2,
            "Existing Provisional Consensus Label": _text(r["provisional_consensus_label"]),
            "Original Adjudication Reason": _text(r["adjudication_reason"]),
            "Disagreement Type": pair,
            "Disagreement Kind": kind,
            "Candidate Categories": " | ".join(c for c in dict.fromkeys(
                candidates + [x for x in (r1, r2) if x != NON_FINAL])),
            "Definition Issues": ", ".join(defs),
            **{f"Field: {f}": v for f, v in relevant.items()},
            "Evidence Cited by Reviewer 1": cited_evidence(r, 1),
            "Evidence Cited by Reviewer 2": cited_evidence(r, 2),
            "Field Conflicts or Gaps": " | ".join(gaps),
            "Evidence Still Needed": question,
            "Adjudication Priority": priority,
            "Priority Reason": reason,
            "Final Adjudicated Label": "",
            "Adjudicator Rationale": "",
            "Adjudicator": "",
            "Adjudication Date": "",
            "Adjudication Status": "PENDING",
        })
    queue = pd.DataFrame(rows)
    order = {"HIGH": 0, "MEDIUM": 1, "LOW": 2}
    queue["_p"] = queue["Adjudication Priority"].map(order)
    queue["_n"] = queue.groupby("Disagreement Type")["Review Case ID"].transform("size")
    queue = queue.sort_values(["_p", "_n", "Disagreement Type", "Review Case ID"],
                              ascending=[True, False, True, True]).drop(columns=["_p", "_n"]).reset_index(drop=True)
    queue.insert(0, "Queue Rank", range(1, len(queue) + 1))

    source_records = disputed.drop(columns=["queue_row"]).rename(columns={"source_row": "adjudication_queue_row_index"})
    source_records.insert(1, "record_id", linked["record_id"].values)
    source_records.insert(2, "source_record_id", linked["source_record_id"].values)
    generator_labels = pd.DataFrame({
        "Review Case ID": disputed["review_case_id"], "Record ID": linked["record_id"],
        "Generator Label (UNVERIFIED_SYNTHETIC_LABEL)": linked["voucher_type"]})

    types = (queue.groupby(["Disagreement Type", "Disagreement Kind", "Adjudication Priority", "Definition Issues"])
             .size().rename("Records").reset_index().sort_values("Records", ascending=False))
    all_reviewed = pd.concat([consensus, src], ignore_index=True)
    return {
        "queue": queue, "source_records": source_records, "generator_labels": generator_labels, "types": types,
        "counts": {
            "reviewed_records": int(len(all_reviewed)),
            "unique_reviewed_ids": int(all_reviewed["review_case_id"].nunique()),
            "consensus_records": int(len(consensus)),
            "needs_adjudication_true_in_all_sheets": int((all_reviewed["needs_adjudication"] == True).sum()),  # noqa: E712
            "summary_sheet": sheets["Summary"].to_dict("records"),
        },
    }


def write_workbook(result, path: Path):
    queue = result["queue"]
    readme = pd.DataFrame([
        ("Purpose", "Prepared queue for qualified human adjudication of the disputed reviewer records. No label has been adjudicated."),
        ("Source of truth", f"{COMPARISON.relative_to(ROOT)} / Adjudication_Queue (needs_adjudication = True). "
                            f"sha256 {hashlib.sha256(COMPARISON.read_bytes()).hexdigest()}"),
        ("Record IDs and fields", f"{DATASET.relative_to(ROOT)} / human_review_queue, matched by review case number; "
                                  f"sha256 {hashlib.sha256(DATASET.read_bytes()).hexdigest()}"),
        ("Columns to fill", "Final Adjudicated Label, Adjudicator Rationale, Adjudicator, Adjudication Date, Adjudication Status "
                            "(yellow). Do not edit any other column."),
        ("Status values", "PENDING (initial) | ADJUDICATED | NEEDS_POLICY_DECISION (blocked by a definition issue) | "
                          "INSUFFICIENT_EVIDENCE (record cannot support any label)"),
        ("Evidence columns", "'Evidence Cited by Reviewer N' quotes the reviewer's own label, confidence, flag, note and cited "
                             "fields; reviewer 1 cited field names, whose values are filled in from the record. 'Field Conflicts "
                             "or Gaps' lists factual contradictions or missing fields; it never proposes a label."),
        ("Priority HIGH", "Two concrete labels conflict on an open definition issue, or both reviewers abstained on a "
                          "record that decides the Other / Miscellaneous policy. Field conflicts are listed per record "
                          "but do not change priority."),
        ("Priority MEDIUM", "Concrete labels conflict without a definition issue; reviewers agree but raised a quality flag; "
                            "one reviewer abstained on an open definition issue; or both abstained on another record."),
        ("Priority LOW", "One reviewer abstained and no open definition issue applies."),
        ("Definition issues", "; ".join(f"{k}: {v}" for k, v in DEFINITION_ISSUES.items())
                              + ". Details: reports/vyom_label_definition_gaps.md"),
        ("document_type caution", "In this synthetic dataset document_type maps 1:1 to the generator's label for 15 classes; "
                                  "reviewer 1 cited it as evidence. Treat it as generated metadata, not independent evidence."),
        ("Generator label", "Sheet Generator_Label_AUDIT_ONLY holds the unverified synthetic label. Hide it from adjudicators "
                            "until they have recorded a decision, to avoid anchoring."),
    ], columns=["Item", "Detail"])
    lists = pd.DataFrame({"Categories": sorted(set(pd.read_excel(DATASET, sheet_name="label_guidelines_v5")["voucher_type"]))})
    status_list = pd.DataFrame({"Statuses": STATUSES})

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        readme.to_excel(writer, sheet_name="README", index=False)
        queue.to_excel(writer, sheet_name="Adjudication_Queue", index=False)
        result["types"].to_excel(writer, sheet_name="Disagreement_Types", index=False)
        pd.DataFrame(list(DEFINITION_ISSUES.items()), columns=["Issue", "Question area"]).to_excel(
            writer, sheet_name="Definition_Issues", index=False)
        result["source_records"].to_excel(writer, sheet_name="Source_Records", index=False)
        result["generator_labels"].to_excel(writer, sheet_name="Generator_Label_AUDIT_ONLY", index=False)
        pd.concat([lists, status_list], axis=1).to_excel(writer, sheet_name="Lists", index=False)

    wb = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="1F3864")
    input_fill = PatternFill("solid", fgColor="FFFF00")
    editable = ["Final Adjudicated Label", "Adjudicator Rationale", "Adjudicator", "Adjudication Date", "Adjudication Status"]
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                cell.font = Font(name="Arial", size=10, bold=cell.row == 1, color="FFFFFF" if cell.row == 1 else "000000")
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for cell in ws[1]:
            cell.fill = header_fill
        for idx, col in enumerate(ws.columns, start=1):
            longest = max(len(str(c.value)) if c.value is not None else 0 for c in col)
            ws.column_dimensions[get_column_letter(idx)].width = max(10, min(longest + 2, 60))
        ws.freeze_panes = "A2"
        if ws.max_row > 1:
            ws.auto_filter.ref = ws.dimensions
    ws = wb["Adjudication_Queue"]
    headers = [c.value for c in ws[1]]
    n = ws.max_row
    for name in editable:
        col = get_column_letter(headers.index(name) + 1)
        for r in range(2, n + 1):
            ws[f"{col}{r}"].fill = input_fill
    label_col = get_column_letter(headers.index("Final Adjudicated Label") + 1)
    status_col = get_column_letter(headers.index("Adjudication Status") + 1)
    dv_label = DataValidation(type="list", formula1=f"=Lists!$A$2:$A${len(lists) + 1}", allow_blank=True)
    dv_status = DataValidation(type="list", formula1=f"=Lists!$B$2:$B${len(STATUSES) + 1}", allow_blank=False)
    ws.add_data_validation(dv_label)
    ws.add_data_validation(dv_status)
    dv_label.add(f"{label_col}2:{label_col}{n}")
    dv_status.add(f"{status_col}2:{status_col}{n}")
    wb["Generator_Label_AUDIT_ONLY"].sheet_state = "hidden"
    wb.save(path)


def main():
    parser = argparse.ArgumentParser(description="Build the human adjudication queue (no labels are decided).")
    parser.add_argument("--output", default=str(OUTPUT))
    args = parser.parse_args()
    out = Path(args.output)
    if out.exists():
        raise FileExistsError(f"Refusing to overwrite {out}")
    result = build_queue()
    write_workbook(result, out)
    q = result["queue"]
    print(f"Saved {out} ({len(q)} records)")
    print(q["Adjudication Priority"].value_counts().to_string())
    print(result["types"].to_string(index=False))
    print(result["counts"])


if __name__ == "__main__":
    main()
